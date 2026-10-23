import logging
import hashlib
import os
from pathlib import Path
from typing import BinaryIO, Iterable, NoReturn
from zipfile import ZipFile

from docx import Document
from docx.document import Document as DocumentType
import httpx
from fastapi import Depends, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from dotenv import load_dotenv
from groq import (
    APIConnectionError,
    APIError,
    APIStatusError,
    APITimeoutError,
    AuthenticationError,
    BadRequestError,
    Groq,
    RateLimitError,
    UnprocessableEntityError,
)
from pydantic import BaseModel, Field, ValidationError
from pypdf import PdfReader

load_dotenv(Path(__file__).with_name(".env"))

logger = logging.getLogger(__name__)
bearer_scheme = HTTPBearer(auto_error=False)


def _configured_positive_limit(name: str, default: int) -> int:
    value = os.getenv(name)
    if value is None:
        return default
    try:
        limit = int(value)
    except ValueError as error:
        raise RuntimeError(f"{name} must be a positive integer.") from error
    if limit <= 0:
        raise RuntimeError(f"{name} must be a positive integer.")
    return limit


MAX_UPLOAD_SIZE_BYTES = _configured_positive_limit("MAX_UPLOAD_SIZE_BYTES", 5 * 1024 * 1024)
MAX_EXTRACTED_TEXT_LENGTH = _configured_positive_limit("MAX_EXTRACTED_TEXT_LENGTH", 100_000)
MAX_JOB_DESCRIPTION_LENGTH = _configured_positive_limit("MAX_JOB_DESCRIPTION_LENGTH", 20_000)
MAX_IMPROVE_TEXT_LENGTH = _configured_positive_limit("MAX_IMPROVE_TEXT_LENGTH", 10_000)
MAX_JOB_ROLE_LENGTH = _configured_positive_limit("MAX_JOB_ROLE_LENGTH", 200)
MAX_DOCX_UNCOMPRESSED_SIZE_BYTES = _configured_positive_limit(
    "MAX_DOCX_UNCOMPRESSED_SIZE_BYTES",
    25 * 1024 * 1024,
)
AI_RATE_LIMIT_WINDOW_SECONDS = _configured_positive_limit("AI_RATE_LIMIT_WINDOW_SECONDS", 60)
AI_RATE_LIMIT_USER_REQUESTS = _configured_positive_limit("AI_RATE_LIMIT_USER_REQUESTS", 10)
AI_RATE_LIMIT_IP_REQUESTS = _configured_positive_limit("AI_RATE_LIMIT_IP_REQUESTS", 30)
UPLOAD_READ_CHUNK_SIZE = 64 * 1024

RATE_LIMIT_SCRIPT = """
local user_count = redis.call('INCR', KEYS[1])
if user_count == 1 then
    redis.call('EXPIRE', KEYS[1], ARGV[1])
end
local ip_count = redis.call('INCR', KEYS[2])
if ip_count == 1 then
    redis.call('EXPIRE', KEYS[2], ARGV[1])
end
local retry_after = math.max(redis.call('TTL', KEYS[1]), redis.call('TTL', KEYS[2]))
if user_count > tonumber(ARGV[2]) or ip_count > tonumber(ARGV[3]) then
    return {1, retry_after}
end
return {0, retry_after}
"""

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
        "https://ai-resume-reviewer1.vercel.app",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.get("/")
def read_root():
    return {"message": "AI Resume Reviewer backend is running!"}


MODEL = "openai/gpt-oss-120b"


class ResumeAnalysis(BaseModel):
    overall_score: int = Field(ge=0, le=100)
    strengths: list[str]
    weaknesses: list[str]
    missing_skills: list[str]
    ats_suggestions: list[str]
    formatting_suggestions: list[str]


class JobMatch(BaseModel):
    match_percentage: int = Field(ge=0, le=100)
    matching_skills: list[str]
    missing_skills: list[str]
    recommended_keywords: list[str]


class ImprovedBullet(BaseModel):
    improved: str


def _safe_error_identifier(value: object) -> str | None:
    if (
        not isinstance(value, str)
        or not value
        or len(value) > 64
        or not value.isascii()
        or not all(character.isalnum() or character in "._-" for character in value)
    ):
        return None
    return value


def _log_groq_api_error(category: str, error: APIError) -> None:
    status_code = None
    request_id = None
    provider_type = None
    provider_code = None

    if isinstance(error, APIStatusError):
        status_code = error.status_code
        request_id = _safe_error_identifier(error.response.headers.get("x-request-id"))

    if isinstance(error.body, dict):
        details = error.body.get("error", error.body)
        if isinstance(details, dict):
            provider_type = _safe_error_identifier(details.get("type"))
            provider_code = _safe_error_identifier(details.get("code"))

    logger.error(
        "Groq request failed: category=%s exception=%s model=%s status=%s "
        "provider_type=%s provider_code=%s request_id=%s",
        category,
        type(error).__name__,
        MODEL,
        status_code,
        provider_type,
        provider_code,
        request_id,
    )


def _raise_ai_service_error(category: str, error: APIError) -> NoReturn:
    _log_groq_api_error(category, error)
    raise HTTPException(
        status_code=502,
        detail="The AI service could not analyze the resume.",
    ) from error


def _authentication_error() -> HTTPException:
    return HTTPException(
        status_code=401,
        detail="Authentication required.",
        headers={"WWW-Authenticate": "Bearer"},
    )


async def require_authenticated_user(
    credentials: HTTPAuthorizationCredentials | None = Depends(bearer_scheme),
) -> str:
    if credentials is None or not credentials.credentials:
        raise _authentication_error()

    supabase_url = os.getenv("SUPABASE_URL")
    supabase_publishable_key = os.getenv("SUPABASE_PUBLISHABLE_KEY")
    if not supabase_url or not supabase_publishable_key:
        logger.error("Supabase authentication verification is not configured.")
        raise HTTPException(
            status_code=503,
            detail="Authentication service is unavailable.",
        )
    try:
        async with httpx.AsyncClient(timeout=httpx.Timeout(5.0, connect=2.0)) as client:
            response = await client.get(
                f"{supabase_url.rstrip('/')}/auth/v1/user",
                headers={
                    "apikey": supabase_publishable_key,
                    "Authorization": f"Bearer {credentials.credentials}",
                },
            )
    except httpx.TimeoutException as error:
        logger.warning(
            "Supabase authentication verification timed out: exception=%s",
            type(error).__name__,
        )
        raise HTTPException(
            status_code=503,
            detail="Authentication service is unavailable.",
        ) from error
    except httpx.RequestError as error:
        logger.warning(
            "Supabase authentication verification failed: exception=%s",
            type(error).__name__,
        )
        raise HTTPException(
            status_code=503,
            detail="Authentication service is unavailable.",
        ) from error

    if response.status_code in (401, 403):
        raise _authentication_error()
    if response.status_code != 200:
        logger.warning(
            "Supabase authentication verification returned HTTP %s.",
            response.status_code,
        )
        raise HTTPException(
            status_code=503,
            detail="Authentication service is unavailable.",
        )
    try:
        verified_user = response.json()
    except ValueError as error:
        logger.warning("Supabase authentication verification returned invalid JSON.")
        raise HTTPException(
            status_code=503,
            detail="Authentication service is unavailable.",
        ) from error

    user_id = verified_user.get("id") if isinstance(verified_user, dict) else None
    if not isinstance(user_id, str) or not user_id:
        logger.warning("Supabase authentication verification returned an invalid user.")
        raise HTTPException(
            status_code=503,
            detail="Authentication service is unavailable.",
        )
    return user_id


async def enforce_ai_rate_limit(
    request: Request,
    user_id: str = Depends(require_authenticated_user),
) -> None:
    redis_url = os.getenv("RATE_LIMIT_REDIS_URL")
    redis_token = os.getenv("RATE_LIMIT_REDIS_TOKEN")
    if not redis_url or not redis_token:
        logger.error("Distributed AI rate limiting is not configured.")
        raise HTTPException(
            status_code=503,
            detail="AI service is temporarily unavailable.",
        )

    if not redis_url.startswith("https://"):
        logger.error("Distributed AI rate limiting requires an HTTPS Redis REST URL.")
        raise HTTPException(
            status_code=503,
            detail="AI service is temporarily unavailable.",
        )

    client_host = request.client.host if request.client else "unknown"
    user_key = hashlib.sha256(user_id.encode("utf-8")).hexdigest()
    ip_key = hashlib.sha256(client_host.encode("utf-8")).hexdigest()
    try:
        async with httpx.AsyncClient(timeout=httpx.Timeout(3.0, connect=2.0)) as client:
            response = await client.post(
                redis_url.rstrip("/"),
                json=[
                    "EVAL",
                    RATE_LIMIT_SCRIPT,
                    "2",
                    f"ai-rate-limit:user:{user_key}",
                    f"ai-rate-limit:ip:{ip_key}",
                    str(AI_RATE_LIMIT_WINDOW_SECONDS),
                    str(AI_RATE_LIMIT_USER_REQUESTS),
                    str(AI_RATE_LIMIT_IP_REQUESTS),
                ],
                headers={"Authorization": f"Bearer {redis_token}"},
            )
    except httpx.RequestError as error:
        logger.warning(
            "Distributed AI rate-limit check failed: exception=%s",
            type(error).__name__,
        )
        raise HTTPException(
            status_code=503,
            detail="AI service is temporarily unavailable.",
        ) from error

    if response.status_code != 200:
        logger.warning(
            "Distributed AI rate-limit store returned HTTP %s.",
            response.status_code,
        )
        raise HTTPException(
            status_code=503,
            detail="AI service is temporarily unavailable.",
        )

    try:
        result = response.json().get("result")
        limited, retry_after = result
        if not isinstance(limited, int) or not isinstance(retry_after, int):
            raise ValueError("Unexpected rate-limit response.")
    except (AttributeError, TypeError, ValueError) as error:
        logger.warning("Distributed AI rate-limit store returned an invalid response.")
        raise HTTPException(
            status_code=503,
            detail="AI service is temporarily unavailable.",
        ) from error

    if limited:
        raise HTTPException(
            status_code=429,
            detail="AI request limit reached. Please try again later.",
            headers={"Retry-After": str(max(1, retry_after))},
        )


async def enforce_upload_size_limit(file: UploadFile) -> None:
    size = 0
    while True:
        chunk = await file.read(min(UPLOAD_READ_CHUNK_SIZE, MAX_UPLOAD_SIZE_BYTES - size + 1))
        if not chunk:
            break
        size += len(chunk)
        if size > MAX_UPLOAD_SIZE_BYTES:
            raise HTTPException(
                status_code=413,
                detail=f"Resume file exceeds the {MAX_UPLOAD_SIZE_BYTES}-byte upload limit.",
            )
    await file.seek(0)


def _join_text_with_limit(pieces: Iterable[str]) -> str:
    parts: list[str] = []
    length = 0
    for piece in pieces:
        if parts:
            length += 1
            if length > MAX_EXTRACTED_TEXT_LENGTH:
                raise HTTPException(
                    status_code=413,
                    detail="The resume contains more readable text than this service can analyze.",
                )
        length += len(piece)
        if length > MAX_EXTRACTED_TEXT_LENGTH:
            raise HTTPException(
                status_code=413,
                detail="The resume contains more readable text than this service can analyze.",
            )
        parts.append(piece)
    return "\n".join(parts)


def _iter_docx_text(document: DocumentType) -> Iterable[str]:
    for paragraph in document.paragraphs:
        yield paragraph.text
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield cell.text


def extract_resume_text(filename: str, content: BinaryIO) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in {".pdf", ".docx"}:
        raise HTTPException(status_code=400, detail="Only PDF and DOCX resumes are supported.")

    try:
        content.seek(0)
        header = content.read(1024)
        if extension == ".pdf":
            if b"%PDF-" not in header:
                raise HTTPException(status_code=400, detail="The uploaded file is not a valid PDF.")
            content.seek(0)
            reader = PdfReader(content)
            if reader.is_encrypted:
                raise HTTPException(
                    status_code=400,
                    detail="Password-protected PDF resumes are not supported.",
                )
            text = _join_text_with_limit(page.extract_text() or "" for page in reader.pages)
        elif extension == ".docx":
            if not header.startswith(b"PK\x03\x04"):
                raise HTTPException(status_code=400, detail="The uploaded file is not a valid DOCX.")
            content.seek(0)
            with ZipFile(content) as archive:
                entries = archive.infolist()
                if sum(entry.file_size for entry in entries) > MAX_DOCX_UNCOMPRESSED_SIZE_BYTES:
                    raise HTTPException(
                        status_code=413,
                        detail="The DOCX file expands beyond the supported processing limit.",
                    )
                if any(entry.flag_bits & 0x1 for entry in entries):
                    raise HTTPException(
                        status_code=400,
                        detail="Encrypted DOCX resumes are not supported.",
                    )
                entry_names = {entry.filename for entry in entries}
                if not {"[Content_Types].xml", "word/document.xml"}.issubset(entry_names):
                    raise HTTPException(status_code=400, detail="The uploaded file is not a valid DOCX.")
            content.seek(0)
            document = Document(content)
            text = _join_text_with_limit(_iter_docx_text(document))
    except HTTPException:
        raise
    except Exception as error:
        logger.warning("Resume document parsing failed: exception=%s", type(error).__name__)
        raise HTTPException(status_code=400, detail="The uploaded resume could not be read.") from error

    if not text.strip():
        raise HTTPException(status_code=400, detail="No readable text was found in the resume.")
    return text


def analyze_with_groq(prompt: str, response_model: type[BaseModel]) -> dict:
    api_key = os.getenv("GROQ_API_KEY")
    if not api_key:
        raise HTTPException(status_code=503, detail="The AI service is not configured.")

    try:
        response = Groq(api_key=api_key).chat.completions.create(
            model=MODEL,
            messages=[
                {
                    "role": "system",
                    "content": (
                        "You are an expert resume reviewer. Return only a valid JSON object "
                        "with exactly the fields and value types requested. Do not use markdown."
                    ),
                },
                {"role": "user", "content": prompt},
            ],
            response_format={"type": "json_object"},
            temperature=0.2,
        )
    except AuthenticationError as error:
        _raise_ai_service_error("authentication failure", error)
    except RateLimitError as error:
        _raise_ai_service_error("rate limit", error)
    except (BadRequestError, UnprocessableEntityError) as error:
        _raise_ai_service_error("invalid request", error)
    except APITimeoutError as error:
        _raise_ai_service_error("timeout", error)
    except APIConnectionError as error:
        _raise_ai_service_error("connection failure", error)
    except APIStatusError as error:
        _raise_ai_service_error("provider HTTP failure", error)
    except APIError as error:
        _raise_ai_service_error("other API failure", error)

    if not response.choices or not response.choices[0].message.content:
        logger.warning("Groq returned an empty completion: model=%s", MODEL)
        raise HTTPException(status_code=502, detail="The AI service returned an empty response.")

    content = response.choices[0].message.content
    try:
        result = response_model.model_validate_json(content or "")
    except ValidationError as error:
        logger.warning(
            "Groq returned invalid structured output: exception=%s validation_errors=%d model=%s",
            type(error).__name__,
            error.error_count(),
            MODEL,
        )
        raise HTTPException(status_code=502, detail="The AI service returned an invalid analysis.") from error
    return result.model_dump()


@app.post("/upload")
async def upload_resume(
    file: UploadFile = File(...),
    job_role: str = Form(...),
    _: None = Depends(enforce_ai_rate_limit),
):
    normalized_job_role = job_role.strip()
    if not normalized_job_role:
        raise HTTPException(status_code=400, detail="Job role is required.")
    if len(normalized_job_role) > MAX_JOB_ROLE_LENGTH:
        raise HTTPException(status_code=413, detail="The target job role is too long.")

    try:
        await enforce_upload_size_limit(file)
        resume_text = extract_resume_text(file.filename or "", file.file)
    finally:
        await file.close()
    analysis = analyze_with_groq(
        "Review this resume for the target role below. Score it from 0 to 100 and "
        "provide concise, actionable lists for strengths, weaknesses, missing_skills, "
        "ats_suggestions, and formatting_suggestions. Return a JSON object with exactly "
        "these keys: overall_score (integer), strengths (array of strings), "
        "weaknesses (array of strings), missing_skills (array of strings), "
        "ats_suggestions (array of strings), formatting_suggestions (array of strings).\n\n"
        f"Target role:\n{normalized_job_role}\n\nResume:\n{resume_text}",
        ResumeAnalysis,
    )
    return {
        "filename": file.filename,
        "job_role": normalized_job_role,
        "text_length": len(resume_text),
        "analysis": analysis,
    }


@app.post("/match")
async def match_resume(
    file: UploadFile = File(...),
    job_description: str = Form(...),
    _: None = Depends(enforce_ai_rate_limit),
):
    normalized_job_description = job_description.strip()
    if not normalized_job_description:
        raise HTTPException(status_code=400, detail="Job description is required.")
    if len(normalized_job_description) > MAX_JOB_DESCRIPTION_LENGTH:
        raise HTTPException(status_code=413, detail="The job description is too long to analyze.")

    try:
        await enforce_upload_size_limit(file)
        resume_text = extract_resume_text(file.filename or "", file.file)
    finally:
        await file.close()
    match = analyze_with_groq(
        "Compare the resume to the job description. Estimate match_percentage from 0 to 100 "
        "and list matching_skills, missing_skills, and recommended_keywords as concise arrays "
        "of strings. Return a JSON object with exactly these keys: match_percentage (integer), "
        "matching_skills (array of strings), missing_skills (array of strings), "
        "recommended_keywords (array of strings).\n\n"
        f"Job description:\n{normalized_job_description}\n\nResume:\n{resume_text}",
        JobMatch,
    )
    return match


@app.post("/improve")
async def improve_text(
    original_text: str = Form(...),
    _: None = Depends(enforce_ai_rate_limit),
):
    normalized_text = original_text.strip()
    if not normalized_text:
        raise HTTPException(status_code=400, detail="Text to improve is required.")
    if len(normalized_text) > MAX_IMPROVE_TEXT_LENGTH:
        raise HTTPException(status_code=413, detail="The text to improve is too long.")

    improved = analyze_with_groq(
        "Rewrite this resume bullet point to be more impactful, using strong action verbs "
        "and quantifiable results where reasonable. Return a JSON object with exactly this "
        'key: "improved" (string).\n\n'
        f"Original:\n{normalized_text}",
        ImprovedBullet,
    )
    return {"original": original_text, **improved}
